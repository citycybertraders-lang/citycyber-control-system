"""
fix_txlog_styles.py — One-time script to restyle existing TRANSACTION LOG rows.

Run once: python fix_txlog_styles.py
Creates a backup before modifying. Safe to run multiple times (idempotent).
"""

import os
import shutil
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "data.xlsx")
TXLOG_SHEET = "🧾 TRANSACTION LOG"
HEADER_ROW  = 4
DATA_START  = 5

# ── Style definitions (exact match of existing styled rows) ──────────
FONT_BASE  = Font(name="Segoe UI", size=9, bold=False)
FONT_BOLD  = Font(name="Segoe UI", size=9, bold=True)
CLR_WHITE  = "F0F6FC"
CLR_DIM    = "8B949E"
CLR_GREEN  = "39D353"
FMT_RUPEE  = '₹#,##0.00'
FMT_INT    = '#,##0'
ALIGN_C    = Alignment(horizontal="center", vertical="center")
ALIGN_L    = Alignment(horizontal="left",   vertical="center")
BORDER_T   = Border(
    left=Side(style="thin", color="30363D"),
    right=Side(style="thin", color="30363D"),
    top=Side(style="thin", color="30363D"),
    bottom=Side(style="thin", color="30363D"),
)

# Column → (font, color_hex, number_format, alignment)
COL_STYLES = {
    2:  (FONT_BASE, CLR_DIM,   None,      ALIGN_L),   # Timestamp
    3:  (FONT_BASE, CLR_WHITE,  None,      ALIGN_L),   # Service
    4:  (FONT_BASE, CLR_WHITE,  FMT_INT,   ALIGN_C),   # Qty
    5:  (FONT_BOLD, CLR_WHITE,  FMT_RUPEE, ALIGN_C),   # Revenue
    6:  (FONT_BOLD, CLR_GREEN,  FMT_RUPEE, ALIGN_C),   # Profit
    7:  (FONT_BASE, CLR_DIM,    FMT_RUPEE, ALIGN_C),   # Cost
    8:  (FONT_BOLD, CLR_WHITE,  None,      ALIGN_C),   # Payment
    9:  (FONT_BASE, CLR_DIM,    None,      ALIGN_C),   # Customer
    10: (FONT_BASE, CLR_DIM,    FMT_RUPEE, ALIGN_C),   # Base Price
    11: (FONT_BASE, CLR_DIM,    FMT_RUPEE, ALIGN_C),   # Final Price
    12: (FONT_BASE, CLR_DIM,    None,      ALIGN_C),   # Override Type
    13: (FONT_BASE, CLR_DIM,    None,      ALIGN_C),   # Override Value
}


def style_cell(cell, col):
    spec = COL_STYLES.get(col)
    if not spec:
        return
    font_base, color_hex, num_fmt, align = spec
    cell.font      = Font(name=font_base.name, size=font_base.size,
                          bold=font_base.bold, color=color_hex)
    cell.alignment = align
    cell.border    = BORDER_T
    if num_fmt:
        cell.number_format = num_fmt


def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: {EXCEL_PATH} not found")
        return

    # Backup
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = f"{EXCEL_PATH}.bak_txlog_{ts}"
    shutil.copy2(EXCEL_PATH, backup)
    print(f"Backup: {backup}")

    wb = load_workbook(EXCEL_PATH)
    if TXLOG_SHEET not in wb.sheetnames:
        print(f"ERROR: Sheet '{TXLOG_SHEET}' not found. Available: {wb.sheetnames}")
        wb.close()
        return

    ws = wb[TXLOG_SHEET]
    max_row = ws.max_row or DATA_START

    styled = 0
    for r in range(DATA_START, max_row + 1):
        # Check if row has data (Col B = timestamp)
        if ws.cell(r, 2).value is None:
            continue
        for col in range(2, 14):
            style_cell(ws.cell(r, col), col)
        styled += 1

    wb.save(EXCEL_PATH)
    wb.close()
    print(f"Done: restyled {styled} rows in '{TXLOG_SHEET}' (rows {DATA_START}–{max_row})")


if __name__ == "__main__":
    main()
